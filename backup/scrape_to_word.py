import os
import json
import time
import random
import io
import cloudscraper
from bs4 import BeautifulSoup
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# --- KONFIGURASI ---
INPUT_JSON = "articles_result.json"
OUTPUT_DIR = "word_outputs"
PROGRESS_FILE = "progress.json"
BATCH_SIZE = 100  # 1 file Word berisi 100 artikel agar file tidak terlalu berat

os.makedirs(OUTPUT_DIR, exist_ok=True)

# Inisialisasi Cloudscraper
scraper = cloudscraper.create_scraper(
    browser={
        'browser': 'chrome',
        'platform': 'windows',
        'desktop': True
    }
)

def load_progress() -> set:
    """Membaca daftar URL yang sudah selesai diproses."""
    if os.path.exists(PROGRESS_FILE):
        with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
            return set(json.load(f))
    return set()

def save_progress(completed_urls: set):
    """Menyimpan daftar URL yang selesai diproses ke checkpoint file."""
    with open(PROGRESS_FILE, "w", encoding="utf-8") as f:
        json.dump(list(completed_urls), f, indent=4)

def fetch_html(url: str, max_retries: int = 3) -> str | None:
    """Mengambil HTML dengan retry dan delay acak."""
    for attempt in range(1, max_retries + 1):
        try:
            time.sleep(random.uniform(1.2, 2.5))
            response = scraper.get(url, timeout=25)
            if response.status_code == 200:
                return response.text
            elif response.status_code in [403, 429, 503]:
                wait = attempt * 4 + random.uniform(2, 4)
                print(f"    ⚠️ Block/Limit ({response.status_code}). Wait {wait:.1f}s...")
                time.sleep(wait)
        except Exception as e:
            time.sleep(attempt * 2)
    return None

def extract_article_data(html: str, url: str) -> dict | None:
    """Mengekstrak elemen artikel murni khusus untuk Search Engine Land berdasarkan HTML asli."""
    
    # 0. Filter URL listing/index/tag
    if any(ignore_path in url.lower() for ignore_path in ["/latest-posts", "/category/", "/tag/", "/author/", "/page/"]):
        return None

    soup = BeautifulSoup(html, "html.parser")
    
    # 1. Judul Artikel
    title_tag = soup.find("h1")
    if not title_tag:
        return None
    title = title_tag.get_text().strip()

    if "latest posts" in title.lower():
        return None

    # 2. Cari Container Artikel Utama (Sesuai HTML: #articleContent atau .body-content)
    article_body = (
        soup.find("div", id="articleContent") or
        soup.find("div", class_="body-content") or
        soup.find("div", class_="bialty-container")
    )
    
    if not article_body:
        return None

    # 3. Author & Kategori
    author_tag = soup.find("div", class_="author-name") or soup.find("a", class_=lambda c: c and "author" in c.lower())
    author = author_tag.get_text().strip() if author_tag else "Search Engine Land Staff"

    category_tag = soup.find("a", class_=lambda c: c and "category" in c.lower())
    category = category_tag.get_text().strip() if category_tag else "General"

    # 4. Ambil Gambar Utama (Featured Image / Gambar pertama di artikel)
    img_url = None
    first_figure = article_body.find("figure", class_=lambda c: c and "wp-block-image" in c)
    if first_figure:
        img_tag = first_figure.find("img")
        if img_tag and img_tag.get("src"):
            img_url = img_tag["src"]

    # 5. Pembersihan Elemen Sampah (Iklan, Bio Penulis, Disclosure, Topics)
    for junk in article_body.find_all(["script", "style", "iframe", "aside", "form", "hr"]):
        junk.decompose()

    for junk_class in ["author-about", "article-disclosure", "ttd-topics-display", "ad-space"]:
        for junk in article_body.find_all("div", class_=lambda c: c and junk_class in c):
            junk.decompose()

    # 6. Ekstraksi Konten Teks
    content_elements = []

    # Buka subhead/ringkasan jika ada
    subhead = soup.find("h2", class_="subhead")
    if subhead and subhead.get_text().strip():
        content_elements.append({"type": "paragraph", "text": f"Summary: {subhead.get_text().strip()}"})

    # Telusuri paragraf, heading, dan daftar poin
    for elem in article_body.find_all(["p", "h2", "h3", "h4", "ul", "ol"]):
        tag_name = elem.name
        text = elem.get_text().strip()
        
        if not text:
            continue
            
        # Abaikan kalimat promosi/disclaimer bawaan
        if any(junk_phrase in text.lower() for junk_phrase in ["search engine land is owned by", "related articles", "subscribe to"]):
            continue

        if tag_name in ["h2", "h3", "h4"]:
            content_elements.append({"type": "heading", "text": text, "level": tag_name})
        elif tag_name in ["ul", "ol"]:
            for li in elem.find_all("li"):
                li_text = li.get_text().strip()
                if li_text:
                    content_elements.append({"type": "bullet", "text": li_text})
        elif tag_name == "p":
            content_elements.append({"type": "paragraph", "text": text})

    # Jika tidak ada paragraf yang terekstrak, abaikan
    if len(content_elements) < 1:
        return None

    return {
        "title": title,
        "author": author,
        "category": category,
        "img_url": img_url,
        "elements": content_elements
    }



def add_article_to_doc(doc: Document, article_data: dict, pub_date: str, url: str):
    """Menuliskan data artikel ke dalam dokumen Word."""
    # Title (H1)
    title_p = doc.add_heading(article_data["title"], level=1)
    title_p.style.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    
    # Meta Info
    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(f"📅 Tanggal: {pub_date}  |  ✍️ Penulis: {article_data['author']}  |  🏷️ Kategori: {article_data['category']}")
    meta_run.font.size = Pt(9.5)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # URL Link
    url_p = doc.add_paragraph()
    url_run = url_p.add_run(f"🔗 Source: {url}")
    url_run.font.size = Pt(9)
    url_run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

    # Featured Image (Jika ada)
    if article_data["img_url"]:
        try:
            img_res = scraper.get(article_data["img_url"], timeout=10)
            if img_res.status_code == 200:
                image_stream = io.BytesIO(img_res.content)
                doc.add_picture(image_stream, width=Inches(5.5))
                last_paragraph = doc.paragraphs[-1]
                last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass # Skip jika gambar gagal terdownload

    # Body Content
    for elem in article_data["elements"]:
        if elem["type"] == "heading":
            level = 2 if elem["level"] == "h2" else 3
            doc.add_heading(elem["text"], level=level)
        elif elem["type"] == "bullet":
            doc.add_paragraph(elem["text"], style='List Bullet')
        elif elem["type"] == "paragraph":
            doc.add_paragraph(elem["text"])

    # Separator Garis Pembatas Artikel
    sep_p = doc.add_paragraph()
    sep_run = sep_p.add_run("―" * 55)
    sep_run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    sep_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph() # Spasi antar artikel

def main():
    print("=== MULAI EXTRACTION ARTIKEL KE WORD (.DOCX) ===")
    
    if not os.path.exists(INPUT_JSON):
        print(f"❌ File '{INPUT_JSON}' tidak ditemukan!")
        return

    with open(INPUT_JSON, "r", encoding="utf-8") as f:
        articles_list = json.load(f)

    completed_urls = load_progress()
    pending_articles = [a for a in articles_list if a["url"] not in completed_urls]

    print(f"Total Artikel   : {len(articles_list)}")
    print(f"Sudah Diproses  : {len(completed_urls)}")
    print(f"Sisa Diproses   : {len(pending_articles)}\n")

    if not pending_articles:
        print("✨ Semua artikel sudah selesai diproses ke file Word!")
        return

    # Hitung index batch file Word
    batch_num = (len(completed_urls) // BATCH_SIZE) + 1
    doc = Document()
    current_batch_count = 0

    for idx, item in enumerate(pending_articles, start=1):
        url = item["url"]
        pub_date = item.get("published_date", "N/A")
        global_idx = len(completed_urls) + 1

        print(f"[{global_idx}/{len(articles_list)}] Fetching: {url}")
        html = fetch_html(url)

        if html:
            article_data = extract_article_data(html, url)
            if article_data:
                add_article_to_doc(doc, article_data, pub_date, url)
                current_batch_count += 1
                completed_urls.add(url)
                
                print(f"  ↳ ✅ Sukses disalin: {article_data['title'][:50]}...")
            else:
                print(f"  ↳ ⚠️ Gagal mengekstrak isi dari {url}")
        else:
            print(f"  ↳ ❌ Gagal membuka {url}")

        # Simpan per BATCH_SIZE (100 artikel = 1 File Word)
        if current_batch_count >= BATCH_SIZE or idx == len(pending_articles):
            doc_filename = os.path.join(OUTPUT_DIR, f"SearchEngineLand_Batch_{batch_num:03d}.docx")
            doc.save(doc_filename)
            save_progress(completed_urls)
            print(f"\n💾 [SAVED] {current_batch_count} artikel tersimpan ke '{doc_filename}'\n")

            # Reset untuk batch berikutnya
            batch_num += 1
            doc = Document()
            current_batch_count = 0

    print("==========================================")
    print("✨ SELESAI! Seluruh artikel berhasil tersimpan di folder 'word_outputs/'.")
    print("==========================================")

if __name__ == "__main__":
    main()